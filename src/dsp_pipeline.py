from __future__ import annotations

from pathlib import Path
from time import perf_counter

import numpy as np
from openpyxl import load_workbook
from scipy import signal

FS_ORIGINAL = 500.0
FS_TARGET = 32.0
UPSAMPLE = 8
DOWNSAMPLE = 125
FEATURE_LOW_HZ = 0.5
FEATURE_HIGH_HZ = 12.5

RESAMPLE_PASSBAND_HZ = FEATURE_HIGH_HZ
RESAMPLE_STOPBAND_HZ = 16.0
RESAMPLE_ATTENUATION_DB = 60.0
RESAMPLE_CUTOFF_HZ = (RESAMPLE_PASSBAND_HZ + RESAMPLE_STOPBAND_HZ) / 2.0
_RESAMPLE_FS_INTERMEDIATE = FS_ORIGINAL * UPSAMPLE
_RESAMPLE_WIDTH_NORM = (
    (RESAMPLE_STOPBAND_HZ - RESAMPLE_PASSBAND_HZ)
    / (_RESAMPLE_FS_INTERMEDIATE / 2.0)
)
RESAMPLE_TAPS, RESAMPLE_KAISER_BETA = signal.kaiserord(
    RESAMPLE_ATTENUATION_DB,
    _RESAMPLE_WIDTH_NORM,
)
if RESAMPLE_TAPS % 2 == 0:
    RESAMPLE_TAPS += 1

DEFAULT_BANDS = {
    "low_0p5_4_hz": (0.5, 4.0),
    "mid_4_8_hz": (4.0, 8.0),
    "high_8_12p5_hz": (8.0, 12.5),
}


def load_eeg_xlsx(path):
    """Load the first worksheet/column as a 1-D EEG vector."""
    wb = load_workbook(Path(path), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    values = [row[0].value for row in ws.iter_rows()]
    wb.close()
    data = np.asarray(values, dtype=float)
    if np.isnan(data).any():
        raise ValueError("EEG input contains missing/non-numeric samples.")
    return data


def design_resample_fir():
    """Design the specification-driven Kaiser anti-alias FIR."""
    return signal.firwin(
        RESAMPLE_TAPS,
        RESAMPLE_CUTOFF_HZ,
        fs=_RESAMPLE_FS_INTERMEDIATE,
        window=("kaiser", RESAMPLE_KAISER_BETA),
    )


def filter_response_metrics(taps=None):
    """Summarise achieved passband and stopband behaviour."""
    taps = design_resample_fir() if taps is None else taps
    freqs, response = signal.freqz(
        taps,
        worN=131072,
        fs=_RESAMPLE_FS_INTERMEDIATE,
    )
    magnitude = np.maximum(np.abs(response), np.finfo(float).tiny)
    db = 20.0 * np.log10(magnitude)
    passband = freqs <= RESAMPLE_PASSBAND_HZ
    stopband = freqs >= RESAMPLE_STOPBAND_HZ
    passband_edge_db = float(np.interp(RESAMPLE_PASSBAND_HZ, freqs, db))
    stopband_edge_db = float(np.interp(RESAMPLE_STOPBAND_HZ, freqs, db))
    return {
        "num_taps": int(len(taps)),
        "kaiser_beta": float(RESAMPLE_KAISER_BETA),
        "passband_edge_hz": RESAMPLE_PASSBAND_HZ,
        "stopband_edge_hz": RESAMPLE_STOPBAND_HZ,
        "passband_edge_gain_db": passband_edge_db,
        "worst_case_passband_deviation_db": float(np.max(np.abs(db[passband]))),
        "stopband_edge_gain_db": stopband_edge_db,
        "minimum_stopband_attenuation_db": float(-np.max(db[stopband])),
    }


def resample_direct_reference(data, taps=None):
    """Direct-form upsample-filter-decimate reference implementation."""
    taps = design_resample_fir() if taps is None else taps
    upsampled = np.zeros(len(data) * UPSAMPLE, dtype=float)
    upsampled[::UPSAMPLE] = data
    filtered = signal.lfilter(taps, 1.0, upsampled) * UPSAMPLE
    delay = (len(taps) - 1) // 2
    return filtered[delay::DOWNSAMPLE]


def resample_polyphase(data, taps=None):
    """Efficient polyphase form using the same FIR coefficients."""
    taps = design_resample_fir() if taps is None else taps
    return signal.resample_poly(
        data,
        UPSAMPLE,
        DOWNSAMPLE,
        window=taps,
        padtype="constant",
    )


def apply_bandpass(data, fs, low=FEATURE_LOW_HZ, high=FEATURE_HIGH_HZ, order=4):
    """Apply a fourth-order Butterworth band-pass with zero-phase filtering."""
    nyquist = fs / 2.0
    b, a = signal.butter(order, [low / nyquist, high / nyquist], btype="band")
    return signal.filtfilt(b, a, data)


def segment_epochs(data, fs, seconds=8.0, overlap=0.5):
    n = int(seconds * fs)
    step = int(n * (1.0 - overlap))
    starts = range(0, len(data) - n + 1, step)
    epochs = np.asarray([data[start : start + n] for start in starts])
    times = np.asarray([(start + n / 2.0) / fs for start in starts])
    return epochs, times


def epoch_psd(
    epochs,
    fs,
    low=FEATURE_LOW_HZ,
    high=FEATURE_HIGH_HZ,
    window="hann",
):
    """One-sided periodogram PSD for each epoch over the requested band."""
    epochs = np.asarray(epochs, dtype=float)
    if epochs.ndim != 2 or len(epochs) == 0:
        raise ValueError("epochs must be a non-empty 2-D array")
    freqs, psd = signal.periodogram(
        epochs,
        fs=fs,
        axis=1,
        window=window,
        detrend=False,
        scaling="density",
    )
    mask = (freqs >= low) & (freqs <= high)
    return freqs[mask], psd[:, mask]


def normalize_psd(psd):
    """Normalize each PSD row to unit total spectral power."""
    psd = np.asarray(psd, dtype=float)
    totals = psd.sum(axis=1, keepdims=True)
    return np.divide(psd, totals, out=np.zeros_like(psd), where=totals != 0)


def iwmf_features(epochs, fs, low=FEATURE_LOW_HZ, high=FEATURE_HIGH_HZ):
    """Intensity-weighted mean frequency using the unwindowed FFT definition."""
    freqs, psd = epoch_psd(epochs, fs, low, high, window="boxcar")
    denominator = psd.sum(axis=1)
    numerator = (psd * freqs).sum(axis=1)
    return np.divide(
        numerator,
        denominator,
        out=np.zeros_like(numerator),
        where=denominator != 0,
    )


def spectral_entropy_features(epochs, fs, low=FEATURE_LOW_HZ, high=FEATURE_HIGH_HZ):
    """Normalized Shannon entropy of the Hann-windowed band-limited PSD."""
    _, psd = epoch_psd(epochs, fs, low, high)
    p = normalize_psd(psd)
    safe = np.where(p > 0, p, 1.0)
    entropy = -np.sum(p * np.log(safe), axis=1)
    return entropy / np.log(p.shape[1])


def spectral_edge_frequency_features(
    epochs,
    fs,
    fraction=0.95,
    low=FEATURE_LOW_HZ,
    high=FEATURE_HIGH_HZ,
):
    """Frequency below which the requested fraction of band power is contained."""
    if not 0.0 < fraction < 1.0:
        raise ValueError("fraction must lie between 0 and 1")
    freqs, psd = epoch_psd(epochs, fs, low, high)
    cumulative = np.cumsum(psd, axis=1)
    totals = cumulative[:, -1]
    targets = totals * fraction
    indices = np.asarray(
        [
            int(np.searchsorted(cumulative[i], targets[i], side="left"))
            for i in range(len(psd))
        ]
    )
    indices = np.clip(indices, 0, len(freqs) - 1)
    return freqs[indices]


def band_power_features(epochs, fs, bands=None):
    """Integrated Hann-windowed PSD for each named frequency band and epoch."""
    bands = DEFAULT_BANDS if bands is None else bands
    max_high = max(high for _, high in bands.values())
    min_low = min(low for low, _ in bands.values())
    freqs, psd = epoch_psd(epochs, fs, min_low, max_high)
    result = {}
    for name, (low, high) in bands.items():
        mask = (freqs >= low) & (freqs <= high)
        result[name] = np.trapezoid(psd[:, mask], freqs[mask], axis=1)
    return result


def comparison_metrics(reference, candidate):
    reference = np.asarray(reference, dtype=float)
    candidate = np.asarray(candidate, dtype=float)
    if len(reference) != len(candidate):
        raise ValueError(f"Length mismatch: {len(reference)} vs {len(candidate)}")
    error = candidate - reference
    mse = float(np.mean(error**2))
    reference_constant = np.allclose(reference, reference[0])
    candidate_constant = np.allclose(candidate, candidate[0])
    if reference_constant or candidate_constant:
        correlation = 1.0 if np.allclose(reference, candidate) else 0.0
    else:
        correlation = float(np.corrcoef(reference, candidate)[0, 1])
    return {
        "n": int(len(reference)),
        "mse": mse,
        "rmse": float(np.sqrt(mse)),
        "correlation": correlation,
    }


def _summary(values):
    values = np.asarray(values, dtype=float)
    return {
        "mean": float(np.mean(values)),
        "median": float(np.median(values)),
        "min": float(np.min(values)),
        "max": float(np.max(values)),
    }


def spectral_preservation_metrics(
    reference_epochs,
    reference_fs,
    candidate_epochs,
    candidate_fs,
    bands=None,
    low=FEATURE_LOW_HZ,
    high=FEATURE_HIGH_HZ,
):
    """Validate target-band spectral shape, band power and summary features."""
    if len(reference_epochs) != len(candidate_epochs):
        raise ValueError("reference and candidate must contain the same number of epochs")

    ref_freqs, ref_psd = epoch_psd(reference_epochs, reference_fs, low, high)
    cand_freqs, cand_psd = epoch_psd(candidate_epochs, candidate_fs, low, high)

    if len(ref_freqs) == len(cand_freqs) and np.allclose(
        ref_freqs,
        cand_freqs,
        atol=1e-12,
    ):
        common_freqs = ref_freqs
        ref_aligned = ref_psd
        cand_aligned = cand_psd
    else:
        common_freqs = ref_freqs[
            (ref_freqs >= cand_freqs[0]) & (ref_freqs <= cand_freqs[-1])
        ]
        ref_mask = np.isin(ref_freqs, common_freqs)
        ref_aligned = ref_psd[:, ref_mask]
        cand_aligned = np.asarray(
            [np.interp(common_freqs, cand_freqs, row) for row in cand_psd]
        )

    ref_norm = normalize_psd(ref_aligned)
    cand_norm = normalize_psd(cand_aligned)
    correlations = np.asarray(
        [
            np.corrcoef(ref_norm[i], cand_norm[i])[0, 1]
            for i in range(len(ref_norm))
        ]
    )
    normalized_rmse = np.sqrt(np.mean((cand_norm - ref_norm) ** 2, axis=1))
    reference_rms = np.sqrt(np.mean(ref_norm**2, axis=1))
    relative_rmse_percent = 100.0 * np.divide(
        normalized_rmse,
        reference_rms,
        out=np.zeros_like(normalized_rmse),
        where=reference_rms != 0,
    )

    bands = DEFAULT_BANDS if bands is None else bands
    ref_band = band_power_features(reference_epochs, reference_fs, bands)
    cand_band = band_power_features(candidate_epochs, candidate_fs, bands)
    band_errors = {}
    for name in bands:
        errors = 100.0 * np.divide(
            np.abs(cand_band[name] - ref_band[name]),
            np.abs(ref_band[name]),
            out=np.zeros_like(ref_band[name]),
            where=np.abs(ref_band[name]) > np.finfo(float).eps,
        )
        band_errors[name] = _summary(errors)

    ref_entropy = spectral_entropy_features(reference_epochs, reference_fs, low, high)
    cand_entropy = spectral_entropy_features(candidate_epochs, candidate_fs, low, high)
    entropy_cmp = comparison_metrics(ref_entropy, cand_entropy)
    entropy_cmp.update(
        {
            "reference_mean": float(np.mean(ref_entropy)),
            "candidate_mean": float(np.mean(cand_entropy)),
            "mean_absolute_error": float(np.mean(np.abs(cand_entropy - ref_entropy))),
        }
    )

    ref_sef95 = spectral_edge_frequency_features(
        reference_epochs,
        reference_fs,
        0.95,
        low,
        high,
    )
    cand_sef95 = spectral_edge_frequency_features(
        candidate_epochs,
        candidate_fs,
        0.95,
        low,
        high,
    )
    sef95_cmp = comparison_metrics(ref_sef95, cand_sef95)
    sef95_cmp.update(
        {
            "reference_mean_hz": float(np.mean(ref_sef95)),
            "candidate_mean_hz": float(np.mean(cand_sef95)),
            "mean_absolute_error_hz": float(
                np.mean(np.abs(cand_sef95 - ref_sef95))
            ),
        }
    )

    return {
        "frequency_bins": int(len(common_freqs)),
        "frequency_resolution_hz": float(np.median(np.diff(common_freqs))),
        "psd_window": "hann",
        "psd_shape": {
            "correlation": _summary(correlations),
            "normalized_psd_rmse": _summary(normalized_rmse),
            "relative_rmse_percent": _summary(relative_rmse_percent),
        },
        "band_power_relative_error_percent": band_errors,
        "spectral_entropy": entropy_cmp,
        "sef95": sef95_cmp,
    }


def theoretical_complexity(n_samples, n_taps=RESAMPLE_TAPS):
    direct_macs = int(n_samples * UPSAMPLE * n_taps)
    polyphase_core_macs = float(n_samples * n_taps / DOWNSAMPLE)
    return {
        "direct_form_macs": direct_macs,
        "polyphase_core_macs_approx": polyphase_core_macs,
        "theoretical_reduction_factor": direct_macs / polyphase_core_macs,
        "direct_form_intermediate_arrays_bytes": int(2 * n_samples * UPSAMPLE * 8),
    }


def _timed_runs(fn, repeats):
    samples = []
    for _ in range(repeats):
        start = perf_counter()
        fn()
        samples.append(perf_counter() - start)
    values = np.asarray(samples)
    return {
        "median_seconds": float(np.median(values)),
        "mean_seconds": float(np.mean(values)),
        "std_seconds": float(np.std(values)),
        "repeats": int(repeats),
    }


def benchmark_resamplers(data, taps=None, direct_repeats=20, poly_repeats=100):
    taps = design_resample_fir() if taps is None else taps
    for _ in range(3):
        resample_direct_reference(data, taps)
        resample_polyphase(data, taps)
    direct_stats = _timed_runs(
        lambda: resample_direct_reference(data, taps),
        direct_repeats,
    )
    poly_stats = _timed_runs(
        lambda: resample_polyphase(data, taps),
        poly_repeats,
    )
    speedup = direct_stats["median_seconds"] / poly_stats["median_seconds"]
    return {
        "direct_form_reference": direct_stats,
        "polyphase_optimized": poly_stats,
        "median_speedup": float(speedup),
    }